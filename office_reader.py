"""Bounded local Office text extraction shared by Gradio and FastAPI.

No macros, embedded objects or external relationships are executed. Modern
formats are parsed directly; legacy PPT requires local LibreOffice conversion.
"""
from dataclasses import dataclass, field
from io import BytesIO
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
from zipfile import ZipFile

OFFICE_EXTENSIONS = {'.docx', '.pptx', '.ppt'}
OFFICE_MIME = {'.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
               '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
               '.ppt': 'application/vnd.ms-powerpoint'}
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_PARTS = 5000
CONVERSION_TIMEOUT = 60


class OfficeError(ValueError):
    pass


@dataclass
class OfficePart:
    number: int
    reference: str
    text: str


@dataclass
class OfficeContent:
    kind: str
    parts: list[OfficePart]
    warnings: list[str] = field(default_factory=list)

    @property
    def text(self):
        return '\n\n'.join(f'[{p.reference}]\n{p.text}' for p in self.parts)


def _package(data, suffix):
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_PARTS or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                raise OfficeError('Office archive expands beyond the safe limit. Split the document.')
            names = {e.filename for e in entries}
            if len(names) != len(entries):
                raise OfficeError('Office archive contains duplicate parts. Re-export it.')
            required = 'word/document.xml' if suffix == '.docx' else 'ppt/presentation.xml'
            if required not in names or '[Content_Types].xml' not in names:
                raise OfficeError('File content does not match its Office extension. Re-export it.')
            for e in entries:
                if e.flag_bits & 1:
                    raise OfficeError('Encrypted Office document: upload an unlocked copy.')
                if e.file_size > 20 * 1024 * 1024 or e.file_size > max(e.compress_size * 250, 1024 * 1024):
                    raise OfficeError('Office archive part exceeds safe decompression limits.')
                if 'vbaproject' in e.filename.lower():
                    raise OfficeError('Macro-containing files are not supported. Export a macro-free copy.')
                if e.filename.endswith(('.xml', '.rels')):
                    xml = archive.read(e)
                    if b'<!DOCTYPE' in xml.upper() or b'<!ENTITY' in xml.upper():
                        raise OfficeError('Office XML contains unsupported entity declarations.')
    except OfficeError:
        raise
    except Exception as exc:
        raise OfficeError('Corrupted or unsupported Office document. Export a fresh copy.') from exc


def libreoffice_path():
    candidates = [os.environ.get('LIBREOFFICE_CMD'), shutil.which('soffice'), shutil.which('libreoffice')]
    if os.name == 'nt':
        for variable in ('ProgramFiles', 'ProgramFiles(x86)'):
            candidates.append(str(Path(os.environ.get(variable, 'C:/Program Files')) / 'LibreOffice/program/soffice.exe'))
    candidates.append('/Applications/LibreOffice.app/Contents/MacOS/soffice')
    return next((str(Path(p).resolve()) for p in candidates if p and Path(p).is_file()), None)


def convert_ppt(data, max_bytes, temporary_root=None):
    if not data.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        raise OfficeError('This is not a valid legacy PPT file. Save it as PPTX in PowerPoint.')
    try:
        import olefile
        with olefile.OleFileIO(BytesIO(data)) as compound:
            if not compound.exists('PowerPoint Document'):
                raise OfficeError('Legacy file is not a PowerPoint presentation or is encrypted.')
            if any(any('vba' in part.lower() or part.lower() == 'macros' for part in stream)
                   for stream in compound.listdir()):
                raise OfficeError('Macro-containing PPT files are not supported. Save a macro-free PPTX.')
    except OfficeError:
        raise
    except ImportError as exc:
        raise OfficeError('Legacy PPT needs olefile. Install the project requirements, or save as PPTX.') from exc
    except Exception as exc:
        raise OfficeError('Corrupted or encrypted legacy PPT. Save an unlocked PPTX copy.') from exc
    executable = libreoffice_path()
    if not executable:
        raise OfficeError('Legacy PPT needs local LibreOffice. Install it or save your presentation as PPTX. Set LIBREOFFICE_CMD if it is not on PATH.')
    with tempfile.TemporaryDirectory(prefix='deadlense-ppt-', dir=temporary_root) as directory:
        root = Path(directory)
        profile = root / 'profile'
        (profile / 'user').mkdir(parents=True)
        (profile / 'user/registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Load"><prop oor:name="UpdateLinkMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>''', encoding='utf-8')
        source = root / 'input.ppt'
        source.write_bytes(data)
        try:
            run = subprocess.run([executable, '-env:UserInstallation=' + profile.as_uri(),
                                  '--headless', '--nologo', '--nodefault', '--norestore',
                                  '--convert-to', 'pptx:Impress MS PowerPoint 2007 XML',
                                  '--outdir', str(root), str(source)], capture_output=True,
                                 timeout=CONVERSION_TIMEOUT,
                                 creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        except subprocess.TimeoutExpired as exc:
            raise OfficeError('Legacy PPT conversion timed out. Save it as PPTX and retry.') from exc
        except OSError as exc:
            raise OfficeError('Could not start LibreOffice. Check LIBREOFFICE_CMD or use PPTX.') from exc
        target = root / 'input.pptx'
        if run.returncode or not target.is_file() or not target.stat().st_size:
            raise OfficeError('LibreOffice could not convert this PPT. Save an unlocked PPTX copy.')
        if target.stat().st_size > max_bytes:
            raise OfficeError('Converted presentation exceeds the upload limit. Split it first.')
        return target.read_bytes()


def _table_text(table, depth=0):
    from docx.table import Table
    if depth > 8:
        raise OfficeError('Nested tables exceed the supported depth.')
    rows = []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            chunks = []
            for block in cell.iter_inner_content():
                chunks.append(_table_text(block, depth + 1) if isinstance(block, Table) else block.text)
            cells.append('\n'.join(c for c in chunks if c.strip()))
        rows.append(' | '.join(cells))
    return '\n'.join(rows)


def _shape_text(shapes, depth=0):
    if depth > 8:
        raise OfficeError('Grouped shapes exceed the supported depth.')
    # Approximate visual reading order; retain slide order exactly.
    chunks = []
    for shape in sorted(shapes, key=lambda s: (s.top or 0, s.left or 0)):
        if hasattr(shape, 'shapes'):
            chunks.extend(_shape_text(shape.shapes, depth + 1))
        elif shape.has_table:
            chunks.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
        elif shape.has_text_frame:
            chunks.append(shape.text_frame.text)
    return chunks


def read_office(filename, data, *, max_bytes=20 * 1024 * 1024, max_chars=90000, max_slides=30, temporary_root=None):
    suffix = Path(filename).suffix.lower()
    if suffix not in OFFICE_EXTENSIONS:
        raise OfficeError('Unsupported Office format. Use DOCX, PPTX or PPT.')
    if not data:
        raise OfficeError('The selected Office document is empty (0 bytes).')
    if len(data) > max_bytes:
        raise OfficeError('Office document exceeds the upload size limit.')
    legacy = suffix == '.ppt'
    if legacy:
        data = convert_ppt(data, max_bytes, temporary_root)
        suffix = '.pptx'
    _package(data, suffix)
    parts = []
    warnings = ['Only selectable body/slide text and tables are read. Text inside images, charts, headers, footnotes and speaker notes is not extracted; export as PDF for visual/OCR review.']
    try:
        if suffix == '.docx':
            from docx import Document
            from docx.table import Table
            document = Document(BytesIO(data))
            paragraph, table = 0, 0
            for block in document.iter_inner_content():
                if isinstance(block, Table):
                    table += 1
                    label, text = f'table {table}', _table_text(block)
                else:
                    paragraph += 1
                    label, text = f'paragraph {paragraph}', block.text
                parts.append(OfficePart(len(parts) + 1, label, text.strip()))
                if len(parts) > MAX_PARTS:
                    raise OfficeError('Too many document blocks. Split the Word document.')
                if sum(len(p.text) for p in parts) > max_chars:
                    raise OfficeError('Document text exceeds the character limit. Split it first.')
            kind = 'docx'
        else:
            from pptx import Presentation
            deck = Presentation(BytesIO(data))
            if len(deck.slides) > max_slides:
                raise OfficeError(f'Presentation exceeds {max_slides} slides. Split it first.')
            for number, slide in enumerate(deck.slides, 1):
                text = '\n'.join(_shape_text(slide.shapes)).strip()
                parts.append(OfficePart(number, f'slide {number}', text))
                if not any(c.isalnum() for c in text):
                    warnings.append(f'Slide {number} has no readable text. It may be blank or image-only.')
            kind = 'ppt' if legacy else 'pptx'
    except OfficeError:
        raise
    except ImportError as exc:
        raise OfficeError('Office reader dependencies are missing. Install requirements-ui.txt and requirements.txt.') from exc
    except Exception as exc:
        raise OfficeError('Cannot read this Office document. It may be corrupted, encrypted or unsupported. Re-export it.') from exc
    content = OfficeContent(kind, parts, warnings)
    if not any(any(c.isalnum() for c in p.text) for p in parts):
        raise OfficeError('No readable text found in this document. For image-only content, export as PDF and use OCR.')
    if len(content.text) > max_chars:
        raise OfficeError('Document text exceeds the character limit. Split it first.')
    return content
