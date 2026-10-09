"""Real Office parsing and HTTP integration; Gemma calls are controlled test mocks."""
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from office_reader import OfficeError, read_office, libreoffice_path
from app.ingest import IngestError, read_upload, document_block
from document_reader import read_document, apply_page_edits, file_info
from document_ui import select_document, process_document, save_page, ReaderSession
from backend_client import BackendError
import test_backend_integration as fixture_server

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / 'tests/fixtures'


def docx_bytes(empty=False):
    document = Document()
    if not empty:
        document.add_paragraph('Scholarship notice')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Deadline'
        table.cell(0, 1).text = '15 November 2026'
        document.add_paragraph('Submit application. Bring Student ID.')
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def pptx_bytes(empty=False):
    deck = Presentation()
    if not empty:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
        box.text = 'Scholarship notice'
        deck.slides.add_slide(deck.slide_layouts[6])  # Intentional blank slide.
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(7), Inches(1)).table
        table.cell(0, 0).text = 'Deadline'
        table.cell(0, 1).text = '15 November 2026'
        slide.shapes.add_textbox(Inches(1), Inches(3), Inches(7), Inches(1)).text = 'Submit application. Bring Student ID.'
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


class OfficeReaderTests(unittest.TestCase):
    def test_docx_paragraph_table_order_and_nested_table(self):
        content = read_office('notice.docx', docx_bytes())
        self.assertEqual([p.reference for p in content.parts], ['paragraph 1', 'table 1', 'paragraph 2'])
        self.assertLess(content.text.index('Scholarship'), content.text.index('Deadline'))
        self.assertLess(content.text.index('Deadline'), content.text.index('Submit'))
        document = Document()
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Outer'
        table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = 'Nested'
        output = BytesIO(); document.save(output)
        self.assertIn('Nested', read_office('nested.docx', output.getvalue()).text)

    def test_pptx_slide_order_table_and_blank_slide(self):
        content = read_office('notice.pptx', pptx_bytes())
        self.assertEqual([p.reference for p in content.parts], ['slide 1', 'slide 2', 'slide 3'])
        self.assertEqual(content.parts[1].text, '')
        self.assertIn('Deadline | 15 November 2026', content.parts[2].text)
        self.assertTrue(any('Slide 2' in w for w in content.warnings))

    def test_empty_no_text_and_corrupt(self):
        for extension, data in [('.docx', b''), ('.pptx', b'not a zip'), ('.docx', docx_bytes(True)),
                                ('.pptx', pptx_bytes(True)), ('.ppt', b'not ole')]:
            with self.subTest(extension=extension, size=len(data)), self.assertRaises(OfficeError):
                read_office('notice' + extension, data)
        blank = Presentation(); blank.slides.add_slide(blank.slide_layouts[6])
        output = BytesIO(); blank.save(output)
        with self.assertRaisesRegex(OfficeError, 'No readable text'):
            read_office('blank.pptx', output.getvalue())

    def test_limits_and_wrong_content(self):
        with self.assertRaises(OfficeError):
            read_office('notice.docx', docx_bytes(), max_bytes=1)
        with self.assertRaises(OfficeError):
            read_office('notice.pptx', pptx_bytes(), max_slides=2)
        with self.assertRaises(OfficeError):
            read_office('notice.docx', docx_bytes(), max_chars=3)
        with self.assertRaises(OfficeError):
            read_office('notice.docx', pptx_bytes())
        with self.assertRaises(IngestError):
            read_upload('notice.docx', 'image/png', docx_bytes())
        with self.assertRaises(IngestError):
            read_upload('notice.exe', 'application/octet-stream', docx_bytes())

    def test_archive_macros_entities_and_expansion_rejected(self):
        for name, data in [('word/vbaProject.bin', b'not executable'), ('word/bad.xml', b'<!DOCTYPE evil><x/>')]:
            output = BytesIO()
            with ZipFile(BytesIO(docx_bytes())) as original, ZipFile(output, 'w', ZIP_DEFLATED) as archive:
                for info in original.infolist():
                    archive.writestr(info.filename, original.read(info))
                archive.writestr(name, data)
            with self.assertRaises(OfficeError):
                read_office('notice.docx', output.getvalue())
        with patch('office_reader.MAX_EXPANDED_BYTES', 10), self.assertRaises(OfficeError):
            read_office('notice.docx', docx_bytes())

    def test_local_reader_and_editable_reference_preservation(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for extension, data in [('.docx', docx_bytes()), ('.pptx', pptx_bytes())]:
                file = root / ('notice' + extension); file.write_bytes(data)
                selected = select_document(str(file), ReaderSession())
                self.assertIn(file.name, selected[1])
                self.assertEqual(selected[2], [])
                result = read_document(file, root / extension[1:])
                self.assertIn('15 November 2026', result['extractedText'])
                self.assertFalse(result['modelInferencePerformed'])
                edited = apply_page_edits(result, [p['text'] + '\nReviewed' for p in result['pages']])
                self.assertIn('table 1' if extension == '.docx' else 'slide 3', edited['extractedText'])
                self.assertTrue(edited['edited'])
                self.assertIn('text', result['extractionMethod'])

    def test_txt_preserved_in_picker_and_reader(self):
        with tempfile.TemporaryDirectory() as folder:
            file = Path(folder) / 'notice.txt'; file.write_text('Deadline 15 November 2026')
            self.assertEqual(file_info(file)['fileType'], 'text/plain')
            self.assertIn('Deadline', read_document(file, Path(folder)/'work')['extractedText'])

    def test_office_worker_to_ui_review_and_edit(self):
        names = ['office_notice.docx', 'office_notice.pptx']
        if libreoffice_path():
            names.append('office_notice.ppt')
        for name in names:
            session = ReaderSession()
            try:
                with self.subTest(name=name):
                    events = list(process_document(str(FIXTURES / name), 'English', False, session))
                    result = events[-1][0]
                    self.assertIn('15 November 2026', result['extractedText'])
                    self.assertEqual(events[-1][2], [])  # Text preview, no invented rendered image.
                    page_number = '1' if name.endswith('.docx') else '3'
                    updated, combined, payload, _, _ = save_page(result, page_number, 'User reviewed this text.')
                    self.assertIn('User reviewed this text.', combined)
                    self.assertFalse(payload['modelInferencePerformed'])
                    self.assertTrue(updated['edited'])
            finally:
                session.cancel()

    def test_legacy_missing_converter_and_real_conversion(self):
        legacy = FIXTURES / 'office_notice.ppt'
        if not legacy.exists():
            self.skipTest('No binary PPT fixture available')
        with patch('office_reader.libreoffice_path', return_value=None), self.assertRaisesRegex(OfficeError, 'local LibreOffice'):
            read_office(legacy.name, legacy.read_bytes())
        if not libreoffice_path():
            self.skipTest('LibreOffice not installed; missing-converter validation passed')
        content = read_office(legacy.name, legacy.read_bytes())
        self.assertEqual(content.kind, 'ppt')
        self.assertEqual(len(content.parts), 3)
        self.assertIn('15 November 2026', content.parts[2].text)


class OfficeHttpTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixture_server.IntegrationTests.setUpClass()
        cls.backend = fixture_server.IntegrationTests.client

    @classmethod
    def tearDownClass(cls):
        fixture_server.IntegrationTests.tearDownClass()

    def test_office_upload_review_confirmation_calendar(self):
        for name in ('office_notice.docx', 'office_notice.pptx', 'office_notice.ppt'):
            if name.endswith('.ppt') and not libreoffice_path():
                continue
            with self.subTest(name=name):
                result = self.backend.analyze(path=FIXTURES / name)
                self.assertEqual(result['source']['kind'], Path(name).suffix[1:])
                evidence = next(e for e in result['evidence'] if e['field'] == 'deadline')
                self.assertTrue(evidence['verified'])
                self.assertEqual(evidence['source_ref'], 'table 1' if name.endswith('.docx') else 'slide 3')
                confirmed = self.backend.confirm(result, result['notice'], 'English')
                self.assertTrue(confirmed['checklist']['steps'])
                self.assertIn(b'DTSTART;VALUE=DATE:20261115', self.backend.calendar_file(confirmed))

    def test_extracted_text_actually_reaches_model(self):
        original = fixture_server.controlled_model
        captured = []
        async def capture(messages, schema, **kwargs):
            captured.extend(m['content'] for m in messages)
            return await original(messages, schema, **kwargs)
        with patch.object(fixture_server.ollama_client, 'structured', side_effect=capture):
            self.backend.analyze(path=FIXTURES / 'office_notice.docx')
        joined = '\n'.join(captured)
        self.assertIn('[table 1]', joined)
        self.assertIn('15 November 2026', joined)

    def test_corrupted_and_empty_no_model_call(self):
        with tempfile.TemporaryDirectory() as folder:
            for extension, data in [('.docx', b'corrupted'), ('.pptx', pptx_bytes(True)), ('.ppt', b'broken')]:
                file = Path(folder) / ('invalid' + extension); file.write_bytes(data)
                with patch.object(fixture_server.ollama_client, 'structured') as inference:
                    with self.assertRaises(BackendError):
                        self.backend.analyze(path=file)
                    inference.assert_not_called()
